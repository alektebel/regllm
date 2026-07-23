"""ReAct verification pipeline for DQC generation (anti-hallucination).

Each rule runs through a fresh-context Reason+Act loop:

1. **Sufficiency check** — ONE stateless agent decides whether the rule can
   be built from the dictionary without inventing anything (all fields
   identified, interpretation unambiguous). The claimed fields are verified
   against the dictionary deterministically; a claim of a non-existent
   field flips the verdict to insufficient. Insufficient rules are flagged
   ``ambigua`` with a justification of what is missing — never generated.
2. **SAS generation** — one agent writes the DQC as SAS PROC SQL
   (``SELECT * FROM <tabla> WHERE <condición de error>``).
3. **Validation** — static: every identifier in the query must exist in the
   dictionary (or be a known keyword/function); dynamic (when a data Excel
   with extracted cases was uploaded): the query is executed against the
   cases in an in-memory SQLite database.
4. **Correction loop** — validation errors are fed back to a fresh
   generation agent, up to ``MAX_ATTEMPTS`` times.

When the cases Excel carries a label column (``DQC_ID``/``ID_DQC``/…)
and the rule carries a previous DQC id (``DQC_X: regla`` or ``[DQC_X]
regla``), the executed query also yields precision/recall against the
historically flagged cases.
"""

from __future__ import annotations

import io
import logging
import os
import re
import sqlite3
from typing import Any

from . import dqc_dictionary as dict_ai

logger = logging.getLogger(__name__)

MAX_ATTEMPTS = 3          # generation + correction rounds per rule
MAX_EXAMPLE_ROWS = 50     # violating cases surfaced to the UI (expandable)
MAX_EXAMPLE_COLS = 8
MAX_FETCH_ROWS = 500      # hard cap when executing a query on the cases

# ── previous DQC id ("DQC_X: regla" | "[DQC_X] regla") ───────────────────────

_PREV_ID_RE = re.compile(
    r"^\s*(?:\[(?P<braced>[A-Za-z][\w.-]*)\]|(?P<bare>DQC[\w.-]*))\s*[:\-–—]?\s+")


def _field_embedder():
    """Tier-1 semantic field filter embedder, only when opted in via
    REGLLM_DQC_SEMANTIC_FIELDS. Returns the shared embedding service (which
    itself degrades to zero vectors when no backend is reachable) or None so
    field selection stays purely lexical by default."""
    if (os.getenv("REGLLM_DQC_SEMANTIC_FIELDS") or "").lower() not in (
            "1", "true", "yes", "on"):
        return None
    try:
        from src.knowledge.embeddings import get_embedding_service
        return get_embedding_service()
    except Exception as exc:  # noqa: BLE001 — embeddings optional
        logger.warning("embedding service unavailable: %s", exc)
        return None


def split_prev_id(line: str) -> tuple[str | None, str]:
    """Extract an optional leading previous-DQC id from a rule line."""
    m = _PREV_ID_RE.match(line)
    if not m:
        return None, line.strip()
    prev = m.group("braced") or m.group("bare")
    rest = line[m.end():].strip()
    return (prev, rest) if rest else (None, line.strip())


# ── 1. sufficiency check ─────────────────────────────────────────────────────

SUFFICIENCY_SYSTEM = """\
Eres un experto en calidad de datos para reporting regulatorio bancario.
Recibes UNA regla DQC en lenguaje natural y el diccionario de campos
disponible. Decide si hay información SUFICIENTE Y NO AMBIGUA para construir
un control automático verificable SIN inventar nada.

Marca suficiente=false (ambigua) si se cumple CUALQUIERA de estas:
- Falta algún campo necesario en el diccionario, o el nombre no coincide
  EXACTAMENTE con un campo existente.
- La interpretación de un campo no es inequívoca (unidades, formato, dominio,
  signo, moneda, periodicidad).
- La regla NO define una condición concreta y comprobable: no queda claro qué
  comparación, umbral, operador o valor determina que una fila INCUMPLE.
- Usa términos vagos o subjetivos ("correcto", "razonable", "coherente",
  "válido", "adecuado", "de calidad") sin un criterio medible.
- Presupone un umbral/límite numérico o una lista de valores permitidos que la
  propia regla no aporta.
- Admite más de una interpretación técnica plausible.

Ante la duda, responde suficiente=false: es preferible marcar como ambigua una
regla dudosa que generar un control incorrecto. Responde suficiente=true SOLO
cuando puedas nombrar los campos EXACTOS y la condición exacta de
incumplimiento.

Responde SOLO JSON:
{"suficiente": true|false,
 "campos": ["<campos del diccionario que usará el control>"],
 "interpretacion": "<cómo interpretas la regla y la condición exacta de incumplimiento>",
 "falta": "<si suficiente=false: qué información falta o por qué es ambigua>"}"""


def check_sufficiency(rule: str, fields: list, client) -> dict:
    """Fresh-context sufficiency agent + deterministic field verification."""
    field_names = {f.name.upper() for f in fields}
    relevant = dict_ai.select_relevant_fields(fields, [rule],
                                              embedder=_field_embedder())
    dict_text, _ = dict_ai.fields_to_text(relevant)
    user = f"REGLA DQC:\n{rule}\n\nDICCIONARIO DE CAMPOS:\n{dict_text}"

    try:
        result = client.chat_json(system=SUFFICIENCY_SYSTEM,
                                  user=user[:dict_ai.PROMPT_CHAR_BUDGET],
                                  max_tokens=1024)
    except Exception as exc:  # noqa: BLE001 — verification is best-effort
        logger.warning("sufficiency agent failed: %s", exc)
        # cannot verify — let generation proceed; static validation still
        # protects against hallucinated fields downstream
        return {"suficiente": True, "campos": [], "interpretacion": "",
                "falta": ""}

    if not isinstance(result, dict):
        return {"suficiente": True, "campos": [], "interpretacion": "",
                "falta": ""}

    campos = [str(c).strip().upper() for c in (result.get("campos") or [])
              if str(c).strip()]
    unknown = [c for c in campos if c not in field_names]
    suficiente = bool(result.get("suficiente", False)) and not unknown
    falta = str(result.get("falta") or "").strip()
    if unknown:
        # the model itself referenced fields that do not exist — that IS the
        # hallucination this pipeline exists to catch
        falta = (f"El control necesitaría campos que no existen en el "
                 f"diccionario: {', '.join(unknown)}."
                 + (f" {falta}" if falta else ""))
    if not suficiente and not falta:
        falta = "El modelo no pudo justificar qué información falta."
    return {
        "suficiente": suficiente,
        "campos": [c for c in campos if c in field_names],
        "interpretacion": str(result.get("interpretacion") or "").strip(),
        "falta": falta if not suficiente else "",
    }


# ── 2. SAS generation (with correction feedback) ─────────────────────────────

SAS_GEN_SYSTEM = """\
Eres un experto en SAS para calidad de datos en reporting regulatorio
bancario. Construye UN control DQC para la regla dada usando SOLO campos
del diccionario proporcionado (nombres EXACTOS, en MAYÚSCULAS).

La consulta debe ser PROC SQL de SAS con esta forma EXACTA:
  SELECT * FROM <tabla> WHERE <condición que identifica las filas que VIOLAN la regla>
Usa solo funciones y operadores ANSI (=, <>, <=, >=, AND, OR, NOT, IN,
BETWEEN, IS NULL, ABS, COALESCE, UPPER, SUBSTR...). No uses PROC, DATA
steps, macros ni formatos SAS.

Responde SOLO con JSON: {"dqcs": [<un único objeto>]}. Esquema del objeto:
{
  "dqc_id": "DQC_<CAMPO>_<NNN>",
  "variable": "<campo principal>",
  "descripcion": "<qué verifica>",
  "tipo": "formula|rango|completitud|consistencia|referencial",
  "severidad": "bloqueante|advertencia|informativo",
  "regla_sql": "<SELECT * FROM tabla WHERE ...>",
  "condicion_error": "<cuándo falla>",
  "campos_entrada": ["campo1", "campo2"],
  "referencia_regulatoria": "<del diccionario o 'Sin referencia en diccionario'>",
  "umbral": "<si aplica>",
  "periodicidad": "mensual",
  "justificacion": "<por qué>"
}"""


def generate_sas(rule: str, fields: list, table_name: str, client,
                 campos: list[str] | None = None,
                 feedback: list[str] | None = None,
                 valores: dict[str, list[str]] | None = None) -> Any:
    """One fresh generation agent; ``feedback`` carries validation errors
    from the previous attempt for the correction loop; ``valores`` carries
    sampled real domain values (value grounding)."""
    hint = [rule] + ([" ".join(campos)] if campos else [])
    relevant = dict_ai.select_relevant_fields(fields, hint,
                                              embedder=_field_embedder())
    dict_text, sent = dict_ai.fields_to_text(relevant)
    user = (
        f"Tabla objetivo: {table_name}\n\n"
        f"DICCIONARIO DE CAMPOS ({sent} campos relevantes de {len(fields)}):\n"
        f"{dict_text}\n\n"
        f"REGLA DQC:\n{rule}"
    )
    if valores:
        user += "\n\n" + values_text(valores)
    if feedback:
        user += (
            "\n\nLA CONSULTA ANTERIOR FALLÓ LA VALIDACIÓN:\n- "
            + "\n- ".join(feedback)
            + "\nCorrige la consulta manteniendo la misma regla."
        )
    return client.chat_json(system=SAS_GEN_SYSTEM,
                            user=user[:dict_ai.PROMPT_CHAR_BUDGET],
                            max_tokens=2048)


# ── optional: value grounding (sample real domain values into the prompt) ───

MAX_GROUNDING_FIELDS = 10
MAX_GROUNDING_VALUES = 8


def sample_values(ctx: "CasesContext", campos: list[str]) -> dict[str, list[str]]:
    """Distinct real values per field from the cases Excel, for grounding the
    generation prompt so domain comparisons use values that actually exist
    (semantic value grounding). Purely local — no LLM involved."""
    headers_lower = {h.lower(): h for h in ctx.headers}
    out: dict[str, list[str]] = {}
    for name in campos[:MAX_GROUNDING_FIELDS]:
        col = headers_lower.get(str(name).lower())
        if not col:
            continue
        try:
            rows = ctx.conn.execute(
                f'SELECT DISTINCT "{col}" FROM "{ctx.table}" '
                f'WHERE "{col}" IS NOT NULL LIMIT ?',
                (MAX_GROUNDING_VALUES,)).fetchall()
        except Exception:  # noqa: BLE001 — grounding is best-effort
            continue
        vals = [str(r[0])[:30] for r in rows
                if r[0] is not None and str(r[0]).strip() != ""]
        if vals:
            out[col] = vals
    return out


def values_text(valores: dict[str, list[str]]) -> str:
    lines = [f"- {campo}: " + ", ".join(f"'{v}'" for v in vals)
             for campo, vals in valores.items()]
    return ("VALORES REALES OBSERVADOS EN LOS DATOS (cuando compares contra "
            "literales de dominio, usa EXACTAMENTE estos valores):\n"
            + "\n".join(lines))


# ── optional: semantic judge (does the query implement THIS rule?) ──────────

JUDGE_SYSTEM = """\
Eres un revisor experto de controles de calidad de datos bancarios. Recibes
una regla DQC en lenguaje natural y la consulta generada para implementarla.
La consulta ya es sintácticamente válida; tu trabajo es juzgar la SEMÁNTICA:
- ¿La condición implementa EXACTAMENTE la regla (umbral, sentido, campos)?
- ¿Selecciona las filas que VIOLAN la regla (no las que la cumplen)?
- ¿Falta algún caso límite evidente (nulos, signo, unidades %)?

Responde SOLO JSON:
{"correcto": true|false, "confianza": 0.0-1.0,
 "motivo": "<explicación breve de por qué es correcta o qué está mal>"}"""


def judge_dqc(rule: str, sql: str, descripcion: str, condicion_error: str,
              ejemplos: list[dict] | None, client) -> dict:
    """LLM-as-judge semantic check (CHESS-style unit tester). A failure to
    judge accepts the query — the judge only ever adds scrutiny."""
    user = (f"REGLA DQC:\n{rule}\n\nCONTROL GENERADO:\n"
            f"descripcion: {descripcion}\n"
            f"condicion_error: {condicion_error}\n"
            f"SQL:\n{sql}")
    if ejemplos:
        import json as _json
        user += ("\n\nEJEMPLOS DE FILAS QUE DEVUELVE SOBRE DATOS REALES:\n"
                 + _json.dumps(ejemplos[:3], ensure_ascii=False))
    try:
        result = client.chat_json(system=JUDGE_SYSTEM,
                                  user=user[:dict_ai.PROMPT_CHAR_BUDGET],
                                  max_tokens=512)
    except Exception as exc:  # noqa: BLE001 — judge unavailable ⇒ accept
        logger.warning("semantic judge failed: %s", exc)
        return {"correcto": True, "confianza": 0.0,
                "motivo": "juez no disponible"}
    if not isinstance(result, dict):
        return {"correcto": True, "confianza": 0.0,
                "motivo": "respuesta de juez no interpretable"}
    try:
        confianza = float(result.get("confianza") or 0.0)
    except (TypeError, ValueError):
        confianza = 0.0
    return {
        "correcto": bool(result.get("correcto", True)),
        "confianza": confianza,
        "motivo": str(result.get("motivo") or "").strip(),
    }


# ── 3a. static validation (schema + field existence) ─────────────────────────

_SQL_KEYWORDS = {
    "select", "from", "where", "and", "or", "not", "null", "is", "in",
    "between", "like", "case", "when", "then", "else", "end", "as", "on",
    "join", "inner", "left", "right", "outer", "group", "by", "having",
    "order", "distinct", "asc", "desc", "union", "all", "exists", "cross",
    "calculated",
}

_STRING_RE = re.compile(r"'[^']*'|\"[^\"]*\"")
_IDENT_RE = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
_FUNC_RE = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)\s*\(")
_ALIAS_RE = re.compile(r"\bas\s+([A-Za-z_][A-Za-z0-9_]*)", re.I)


def static_validate(sql: str, fields: list, table_name: str) -> list[str]:
    """Deterministic checks: it's a SELECT and every identifier is a known
    field, keyword, function call, table part, or alias."""
    if not sql or not sql.strip():
        return ["la consulta está vacía"]
    stripped = _STRING_RE.sub("''", sql)
    if not re.search(r"\bselect\b", stripped, re.I):
        return ["la consulta no es un SELECT"]

    field_names = {f.name.lower() for f in fields}
    funcs = {m.lower() for m in _FUNC_RE.findall(stripped)}
    aliases = {m.lower() for m in _ALIAS_RE.findall(stripped)}
    table_parts = {p.lower() for p in re.split(r"\W+", table_name) if p}

    unknown: list[str] = []
    for token in _IDENT_RE.findall(stripped):
        t = token.lower()
        if (t in _SQL_KEYWORDS or t in field_names or t in funcs
                or t in aliases or t in table_parts or len(t) <= 2):
            continue
        if token not in unknown:
            unknown.append(token)
    if unknown:
        return ["la consulta usa campos que NO existen en el diccionario: "
                + ", ".join(unknown[:8])]
    return []


# ── 3b. cases Excel → in-memory SQLite + execution ───────────────────────────

_LABEL_HEADERS = {"dqc_id", "id_dqc", "dqc", "dqc_ids", "dqc_previo",
                  "prev_dqc", "dqc_prev"}
_CASE_COL = "_CASO_"


class CasesContext:
    """Uploaded cases Excel loaded into an in-memory SQLite table."""

    def __init__(self, headers: list[str], rows: list[tuple],
                 label_idx: int | None, table_name: str):
        self.headers = headers
        self.n_rows = len(rows)
        self.label_idx = label_idx
        self.table = re.sub(r"\W+", "_", table_name.split(".")[-1]) or "casos"
        # caso id -> set of previous DQC ids that flagged it historically
        self.labels: dict[int, set[str]] = {}
        if label_idx is not None:
            for i, row in enumerate(rows, start=1):
                value = row[label_idx] if label_idx < len(row) else None
                if value is None or str(value).strip() == "":
                    continue
                ids = {p.strip().upper()
                       for p in re.split(r"[;,|]", str(value)) if p.strip()}
                if ids:
                    self.labels[i] = ids

        # the connection is built on the request's event-loop thread but the
        # SSE generator that queries it runs in a worker thread; access is
        # strictly sequential, so cross-thread use is safe
        self.conn = sqlite3.connect(":memory:", check_same_thread=False)
        cols = ", ".join([f'"{_CASE_COL}" INTEGER']
                         + [f'"{h}"' for h in headers])
        self.conn.execute(f'CREATE TABLE "{self.table}" ({cols})')
        placeholders = ", ".join(["?"] * (len(headers) + 1))
        for i, row in enumerate(rows, start=1):
            values = [i] + [_coerce(row[j]) if j < len(row) else None
                            for j in range(len(headers))]
            self.conn.execute(
                f'INSERT INTO "{self.table}" VALUES ({placeholders})', values)


def _coerce(value: Any) -> Any:
    if value is None or isinstance(value, (int, float)):
        return value
    return str(value)


def _dedupe_headers(headers: list[str], reserved: str) -> list[str]:
    """Make column names non-empty and unique. SQLite is case-insensitive and
    rejects duplicate or blank column names, so an Excel with an empty trailing
    column (two blank headers) or two columns differing only in case fails with
    'duplicate column name'. Real names are kept; only blanks and collisions are
    renamed (blank → columna_N, collision → name_2), so the generated SQL still
    matches its intended columns. Row order/length is preserved."""
    seen = {reserved.lower()}
    out: list[str] = []
    for i, h in enumerate(headers, start=1):
        name = (h or "").strip() or f"columna_{i}"
        base, key, n = name, name.lower(), 1
        while key in seen:
            n += 1
            name = f"{base}_{n}"
            key = name.lower()
        seen.add(key)
        out.append(name)
    return out


def load_cases(data: bytes, table_name: str) -> CasesContext:
    """Parse the cases Excel (first sheet; first non-empty row = headers;
    optional DQC_ID label column) into a queryable context."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                data_only=True)
    ws = wb.active
    headers: list[str] = []
    rows: list[tuple] = []
    if ws is not None:
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None or str(v).strip() == ""
                                  for v in row):
                continue
            if not headers:
                headers = [str(v or "").strip() for v in row]
                continue
            rows.append(row)
    wb.close()
    if not headers:
        raise ValueError("El Excel de datos no tiene cabeceras")
    # find the label column on the ORIGINAL names, then sanitise for SQLite
    label_idx = next((i for i, h in enumerate(headers)
                      if h.lower() in _LABEL_HEADERS), None)
    headers = _dedupe_headers(headers, _CASE_COL)
    return CasesContext(headers, rows, label_idx, table_name)


def _sas_to_sqlite(q: str) -> str:
    """Best-effort rewrite of common SAS PROC SQL idioms to their SQLite
    equivalents, so a query that is valid SAS still runs when we validate it
    against the cases Excel (loaded into SQLite). Conservative — only
    well-known, unambiguous substitutions; the stored/displayed query stays
    SAS, this only affects the in-memory validation run."""
    # SAS not-equal operators → ANSI
    q = q.replace("^=", "<>").replace("~=", "<>")
    # comparison against the SAS numeric-missing dot:  x = .  /  x <> .
    q = re.sub(r"(<>|=)\s*\.(?!\d)",
               lambda m: " IS NOT NULL" if m.group(1) == "<>" else " IS NULL", q)
    # MISSING(expr) → (expr IS NULL)   (one level of parentheses)
    q = re.sub(r"(?i)\bMISSING\s*\(([^()]*)\)", r"(\1 IS NULL)", q)
    # CALCULATED alias reference (SAS-only keyword; SQLite doesn't need it)
    q = re.sub(r"(?i)\bCALCULATED\s+", "", q)
    # function-name differences
    q = re.sub(r"(?i)\bUPCASE\s*\(", "UPPER(", q)
    q = re.sub(r"(?i)\bLOWCASE\s*\(", "LOWER(", q)
    q = re.sub(r"(?i)\bSTRIP\s*\(", "TRIM(", q)
    q = re.sub(r"(?i)\bINDEX\s*\(", "INSTR(", q)
    return q


def run_query(ctx: CasesContext, sql: str, table_name: str) -> dict:
    """Execute a generated query against the cases; returns
    {ok, error?, columnas, ejemplos, n_casos, casos (set of _CASO_ ids)}."""
    q = sql.strip().rstrip(";")
    # tolerate PROC SQL wrappers the model might add despite instructions
    q = re.sub(r"(?im)^\s*(proc\s+sql\s*;|quit\s*;?)\s*$", "", q).strip()
    q = _sas_to_sqlite(q)
    q = re.sub(re.escape(table_name), f'"{ctx.table}"', q, flags=re.I)
    # a bare lib.table nobody declared → last part only
    q = re.sub(r"\b[A-Za-z_]\w*\." + re.escape(ctx.table) + r"\b",
               f'"{ctx.table}"', q, flags=re.I)
    try:
        cur = ctx.conn.execute(q)
        rows = cur.fetchmany(MAX_FETCH_ROWS)
        columns = [d[0] for d in cur.description or []]
    except Exception as exc:  # noqa: BLE001 — execution error feeds the loop
        return {"ok": False, "error": f"la consulta falla al ejecutarse "
                                      f"sobre los casos: {exc}"}

    casos: set[int] = set()
    if _CASE_COL in columns:
        idx = columns.index(_CASE_COL)
        casos = {r[idx] for r in rows if isinstance(r[idx], int)}

    shown_cols = [c for c in columns if c != _CASE_COL][:MAX_EXAMPLE_COLS]
    ejemplos = []
    for r in rows[:MAX_EXAMPLE_ROWS]:
        item = {}
        for c in shown_cols:
            v = r[columns.index(c)]
            item[c] = "" if v is None else str(v)[:40]
        ejemplos.append(item)

    return {"ok": True, "columnas": shown_cols, "ejemplos": ejemplos,
            "n_casos": len(casos) if casos else len(rows), "casos": casos}


def metrics(ctx: CasesContext, casos: set[int],
            prev_id: str | None) -> dict | None:
    """Precision/recall of the flagged cases vs the historical label column.
    Requires both a previous id on the rule and a label column in the data."""
    if not prev_id or ctx.label_idx is None or not casos:
        return None
    pid = prev_id.upper()
    expected = {caso for caso, ids in ctx.labels.items() if pid in ids}
    if not expected:
        return None
    tp = len(casos & expected)
    return {
        "precision": round(tp / len(casos), 3) if casos else 0.0,
        "recall": round(tp / len(expected), 3),
        "esperados": len(expected),
        "aciertos": tp,
    }
